"""
Service d'extraction de texte depuis les ressources pédagogiques Moodle.
Formats supportés : PDF (natif), DOCX, PPTX, TXT, Vidéo (Whisper)
"""

import os
import pdfplumber
from docx import Document
from pptx import Presentation
import yt_dlp


def extract_text(file_path: str) -> dict:
    """
    Point d'entrée unique du service d'extraction.
    Détecte automatiquement le format du fichier via son extension
    et délègue à l'extracteur approprié.

    Args:
        file_path (str): Chemin absolu vers le fichier à extraire.

    Returns:
        dict: {
            "success" (bool)       : True si l'extraction a réussi,
            "text"    (str)        : Texte extrait (vide si échec),
            "page_count" (int|None): Nombre de pages/slides (None si non applicable),
            "format"  (str)        : Format détecté (pdf, docx, pptx, txt, video),
            "error"   (str|None)   : Message d'erreur si échec, None sinon
        }

    Raises:
        Aucune exception levée — les erreurs sont capturées et retournées dans le dict.

    Example:
        >>> result = extract_text("/tmp/cours.pdf")
        >>> if result["success"]:
        ...     print(result["text"])
    """
    if not os.path.exists(file_path):
        return _error(f"Fichier introuvable : {file_path}")

    extension = os.path.splitext(file_path)[1].lower()

    extractors = {
        ".pdf":  _extract_pdf,
        ".docx": _extract_docx,
        ".pptx": _extract_pptx,
        ".txt":  _extract_txt,
        ".mp4":  _extract_video,
        ".avi":  _extract_video,
        ".mov":  _extract_video,
        ".mkv":  _extract_video,
        ".webm": _extract_video,
        ".mp3":  _extract_video,
    }

    if extension not in extractors:
        return _error(
            f"Format non supporté : {extension}. "
            f"Formats acceptés : {', '.join(extractors.keys())}"
        )

    return extractors[extension](file_path)


def extract_text_from_url(url: str, tmp_dir: str) -> dict:
    """
    Télécharge une vidéo depuis une URL externe (YouTube, Vimeo, etc.)
    via yt-dlp, convertit l'audio en MP3 via ffmpeg, puis transcrit
    avec Whisper.

    Args:
        url (str)    : URL de la vidéo externe (YouTube, Vimeo, etc.).
        tmp_dir (str): Répertoire temporaire pour stocker le fichier téléchargé.

    Returns:
        dict: Même structure que extract_text() — texte transcrit avec timestamps.

    Example:
        >>> result = extract_text_from_url("https://youtube.com/watch?v=xxx", "/tmp")
        >>> print(result["text"])
        [00:00] Bonjour et bienvenue dans ce cours...
    """
    try:
        output_path = os.path.join(tmp_dir, "%(id)s.%(ext)s")
        ydl_opts = {
            "format": "bestaudio/best",
            "outtmpl": output_path,
            "quiet": True,
            "postprocessors": [{
                "key": "FFmpegExtractAudio",
                "preferredcodec": "mp3",
                "preferredquality": "192",
            }],
        }
        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            info = ydl.extract_info(url, download=True)
            downloaded_file = ydl.prepare_filename(info)

        # Le fichier est converti en mp3 par le postprocessor
        downloaded_file = downloaded_file.rsplit(".", 1)[0] + ".mp3"

        return _extract_video(downloaded_file)

    except Exception as e:
        return _error(f"Erreur téléchargement URL : {str(e)}", format="video")


# ─────────────────────────────────────────
# Extracteurs par format
# ─────────────────────────────────────────

def _extract_pdf(file_path: str) -> dict:
    """
    Extrait le texte d'un fichier PDF page par page via pdfplumber.
    Détecte automatiquement les PDFs scannés (> 80% de pages sans texte)
    et retourne une erreur explicite dans ce cas.

    Args:
        file_path (str): Chemin absolu vers le fichier PDF.

    Returns:
        dict: Texte extrait avec le nombre de pages, ou erreur si PDF scanné/vide.
    """
    try:
        pages_text = []
        empty_pages = 0

        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)

            for page in pdf.pages:
                text = page.extract_text()
                if text and text.strip():
                    pages_text.append(text.strip())
                else:
                    empty_pages += 1

        if total_pages > 0 and (empty_pages / total_pages) > 0.8:
            return _error(
                "PDF scanné détecté — extraction impossible sans OCR. "
                "Fournissez un PDF avec texte natif.",
                format="pdf",
                page_count=total_pages
            )

        full_text = "\n\n".join(pages_text)

        if not full_text.strip():
            return _error("Aucun texte extractible dans ce PDF", format="pdf")

        return {
            "success": True,
            "text": full_text,
            "page_count": total_pages,
            "format": "pdf",
            "error": None
        }

    except Exception as e:
        return _error(f"Erreur PDF : {str(e)}", format="pdf")


def _extract_docx(file_path: str) -> dict:
    """
    Extrait le texte d'un fichier Word (.docx) paragraphe par paragraphe
    via python-docx. Les paragraphes vides sont ignorés.

    Args:
        file_path (str): Chemin absolu vers le fichier DOCX.

    Returns:
        dict: Texte extrait (page_count = None, DOCX n'a pas de notion de page).
    """
    try:
        doc = Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)

        if not full_text:
            return _error("Aucun texte trouvé dans ce fichier Word", format="docx")

        return {
            "success": True,
            "text": full_text,
            "page_count": None,
            "format": "docx",
            "error": None
        }

    except Exception as e:
        return _error(f"Erreur DOCX : {str(e)}", format="docx")


def _extract_pptx(file_path: str) -> dict:
    """
    Extrait le texte d'un fichier PowerPoint (.pptx) slide par slide
    via python-pptx. Chaque slide est préfixée par [Slide N] pour
    conserver la structure lors du chunking.

    Args:
        file_path (str): Chemin absolu vers le fichier PPTX.

    Returns:
        dict: Texte extrait avec le nombre de slides comme page_count.
    """
    try:
        prs = Presentation(file_path)
        slides_text = []

        for i, slide in enumerate(prs.slides, start=1):
            slide_content = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())

            if slide_content:
                slides_text.append(f"[Slide {i}]\n" + "\n".join(slide_content))

        full_text = "\n\n".join(slides_text)

        if not full_text:
            return _error("Aucun texte trouvé dans ce fichier PowerPoint", format="pptx")

        return {
            "success": True,
            "text": full_text,
            "page_count": len(prs.slides),
            "format": "pptx",
            "error": None
        }

    except Exception as e:
        return _error(f"Erreur PPTX : {str(e)}", format="pptx")


def _extract_txt(file_path: str) -> dict:
    """
    Lit un fichier texte brut (.txt) en UTF-8.
    Bascule automatiquement sur l'encodage latin-1 si UTF-8 échoue
    (utile pour les fichiers générés sur Windows).

    Args:
        file_path (str): Chemin absolu vers le fichier TXT.

    Returns:
        dict: Contenu brut du fichier texte.
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read().strip()

        if not text:
            return _error("Fichier texte vide", format="txt")

        return {
            "success": True,
            "text": text,
            "page_count": None,
            "format": "txt",
            "error": None
        }

    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="latin-1") as f:
                text = f.read().strip()
            return {
                "success": True,
                "text": text,
                "page_count": None,
                "format": "txt",
                "error": None
            }
        except Exception as e:
            return _error(f"Erreur encodage TXT : {str(e)}", format="txt")

    except Exception as e:
        return _error(f"Erreur TXT : {str(e)}", format="txt")


def _extract_video(file_path: str) -> dict:
    """
    Transcrit une vidéo de cours en texte avec timestamps via Whisper (modèle base).
    Chaque segment est formaté [MM:SS] texte pour permettre au chunker
    de conserver la référence temporelle dans chaque chunk.

    Modèle utilisé : base (léger, ~74M paramètres, CPU-only compatible).
    Temps de traitement estimé : ~2-3 min pour une vidéo de 10 min sur CPU.

    Args:
        file_path (str): Chemin absolu vers le fichier vidéo/audio
                         (mp4, avi, mov, mkv, webm, mp3).

    Returns:
        dict: Texte transcrit avec timestamps [MM:SS], page_count = None.

    Example output (text):
        [00:00] Bonjour et bienvenue dans ce cours sur les bases de données.
        [00:15] Aujourd'hui nous allons voir les jointures SQL.
        [01:02] La jointure interne retourne les lignes communes aux deux tables.
    """
    import whisper

    # Assurer que ffmpeg est dans le PATH
    ffmpeg_path = r"C:\Users\wiki\Downloads\ffmpeg-9.0.1-essentials_build\ffmpeg-9.0.1-essentials_build\bin"
    if ffmpeg_path not in os.environ.get("PATH", ""):
        os.environ["PATH"] = ffmpeg_path + os.pathsep + os.environ.get("PATH", "")

    try:
        model = whisper.load_model("base")
        result = model.transcribe(file_path)

        segments = result.get("segments", [])

        if not segments:
            return _error("Aucun segment audio détecté dans la vidéo", format="video")

        lines = []
        for segment in segments:
            start = int(segment["start"])
            minutes = start // 60
            seconds = start % 60
            timestamp = f"[{minutes:02d}:{seconds:02d}]"
            lines.append(f"{timestamp} {segment['text'].strip()}")

        full_text = "\n".join(lines)

        return {
            "success": True,
            "text": full_text,
            "page_count": None,
            "format": "video",
            "error": None
        }

    except Exception as e:
        return _error(f"Erreur transcription vidéo : {str(e)}", format="video")


# ─────────────────────────────────────────
# Helper
# ─────────────────────────────────────────

def _error(message: str, format: str = "unknown", page_count: int = None) -> dict:
    """
    Construit un dictionnaire d'erreur standardisé retourné par tous les extracteurs.

    Args:
        message (str)   : Description de l'erreur.
        format (str)    : Format du fichier concerné (pdf, docx, etc.).
        page_count (int): Nombre de pages si connu malgré l'erreur, None sinon.

    Returns:
        dict: {"success": False, "text": "", "page_count": ..., "format": ..., "error": ...}
    """
    return {
        "success": False,
        "text": "",
        "page_count": page_count,
        "format": format,
        "error": message
    }